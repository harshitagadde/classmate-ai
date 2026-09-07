import os
import io
import time
import datetime
import json
import streamlit as st
import streamlit.components.v1 as components
from dotenv import load_dotenv
from google import genai
from google.genai import types
import pypdf
from pptx import Presentation
import pandas as pd
from PIL import Image

# Load environment variables
load_dotenv()

# Page Setup
st.set_page_config(page_title="Classmate AI", page_icon="📚", layout="wide")

# Persistent File Storage Paths
STUDENT_DB_FILE = "students.json"
ACTIVITY_LOG_FILE = "activity_logs.json"

def load_json_data(file_path, default_data):
    if os.path.exists(file_path):
        try:
            with open(file_path, "r") as f:
                return json.load(f)
        except Exception:
            return default_data
    return default_data

def save_json_data(file_path, data):
    try:
        with open(file_path, "w") as f:
            json.dump(data, f, indent=4)
    except Exception as e:
        st.error(f"Error saving persistent data: {e}")

# Track Session Metrics & Activity Logs Real-Time
if "start_time" not in st.session_state:
    st.session_state.start_time = time.time()
if "notes_processed" not in st.session_state:
    st.session_state.notes_processed = 0
if "quizzes_generated" not in st.session_state:
    st.session_state.quizzes_generated = 0
if "concepts_explained" not in st.session_state:
    st.session_state.concepts_explained = 0
if "messages" not in st.session_state:
    st.session_state.messages = []
if "exam_schedule" not in st.session_state:
    st.session_state.exam_schedule = []
if "todo_list" not in st.session_state:
    st.session_state.todo_list = [
        {"task": "Revise Chapter 1 Notes", "done": False},
        {"task": "Solve Practice Quiz on Math", "done": True}
    ]

# Persistent Student Database Initialization
if "student_db" not in st.session_state:
    st.session_state.student_db = load_json_data(STUDENT_DB_FILE, {
        "harshita@student.com": {"password": "123", "name": "Harshita Gadde"}
    })

# Persistent Activity Log Database Initialization
if "user_activity_log" not in st.session_state:
    st.session_state.user_activity_log = load_json_data(ACTIVITY_LOG_FILE, [])

# Local IST (+5:30) Activity Logger Helper
def log_activity(user_name, user_email, action, details=""):
    utc_now = datetime.datetime.now(datetime.timezone.utc)
    ist_now = utc_now + datetime.timedelta(hours=5, minutes=30)
    timestamp = ist_now.strftime("%Y-%m-%d %I:%M:%S %p IST")
    
    log_entry = {
        "Timestamp": timestamp,
        "Student Name": user_name,
        "Email": user_email,
        "Action Performed": action,
        "Details": details
    }
    
    st.session_state.user_activity_log.insert(0, log_entry)
    save_json_data(ACTIVITY_LOG_FILE, st.session_state.user_activity_log)

# Primary Admin Credentials
if "admin_db" not in st.session_state:
    st.session_state.admin_db = {
        "harshitagadde04@gmail.com": {
            "password": "Honey@peppa2008",
            "name": "Harshita Gadde (Primary Admin)"
        }
    }

if "logged_in_user" not in st.session_state:
    st.session_state.logged_in_user = None
if "user_role" not in st.session_state:
    st.session_state.user_role = "Student"

# API Query Function with Dynamic Model Discovery & 503 Retry Logic
def query_gemini(contents):
    api_key = None
    if "GEMINI_API_KEY" in st.secrets:
        api_key = st.secrets["GEMINI_API_KEY"]
    else:
        api_key = os.getenv("GEMINI_API_KEY")
        
    if not api_key:
        raise Exception("API Key missing! Please add GEMINI_API_KEY in Streamlit Cloud Secrets.")
    
    client = genai.Client(api_key=api_key)
    candidate_models = ['gemini-3.6-flash']
    
    try:
        available_models = [m.name.replace('models/', '') for m in client.models.list() if 'generateContent' in m.supported_generation_methods]
        for m in available_models:
            if m not in candidate_models:
                candidate_models.append(m)
    except Exception:
        pass

    last_err = None
    for model_name in candidate_models:
        for attempt in range(3):
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=contents
                )
                return response.text
            except Exception as e:
                last_err = e
                if "503" in str(e) or "UNAVAILABLE" in str(e):
                    time.sleep(2)
                    continue
                else:
                    break
            
    raise Exception(f"Gemini API Error: {str(last_err)}")

# ---------------------------------------------------------
# SIDEBAR NAVIGATION
# ---------------------------------------------------------
with st.sidebar:
    st.markdown("# 📚 Classmate AI")
    st.caption("Universal Learning & Admin Console")
    st.divider()

    if st.session_state.logged_in_user is not None:
        user_info = st.session_state.logged_in_user
        clean_name = user_info['name'].strip()
        st.success(f"Logged in: **{clean_name}** ({st.session_state.user_role})")
        if st.button("Sign Out", use_container_width=True):
            log_activity(clean_name, st.session_state.logged_in_email, "User Sign Out", "User signed out of portal")
            st.session_state.logged_in_user = None
            st.session_state.logged_in_email = None
            st.rerun()
            
        st.divider()
        if st.session_state.user_role == "Student":
            st.markdown("### 📌 Student Navigation")
            nav = st.radio(
                "Go To:",
                ["🎓 AI Study Assistant", "📋 Custom Study Plan", "💬 Buddy (AI Chatbot)", "📅 Exam Schedule", "📊 Real Study Analytics"]
            )

# ---------------------------------------------------------
# LANDING PAGE WITH DIRECT MOTION SIGN-IN OVERLAY
# ---------------------------------------------------------
if st.session_state.logged_in_user is None:
    
    # Inject background video overlay and custom CSS overrides
    st.markdown("""
        <style>
        /* Zero out top margins and padding */
        .block-container {
            padding-top: 1rem !important;
            padding-bottom: 0rem !important;
        }
        .stApp {
            background-color: transparent !important;
        }
        #bg-video {
            position: fixed;
            right: 0;
            bottom: 0;
            min-width: 100%;
            min-height: 100%;
            z-index: -1;
            object-fit: cover;
            filter: brightness(0.65) contrast(1.1);
        }
        
        /* Glassmorphic Container for Login Panel */
        .motion-login-card {
            background: rgba(24, 24, 27, 0.72) !important;
            backdrop-filter: blur(20px) saturate(120%) !important;
            -webkit-backdrop-filter: blur(20px) saturate(120%) !important;
            border: 1px solid rgba(255, 255, 255, 0.12) !important;
            border-radius: 24px !important;
            padding: 20px 30px 30px 30px !important;
            box-shadow: 0 20px 50px rgba(0, 0, 0, 0.5) !important;
            margin-top: -30px !important;
        }

        .hero-title {
            font-size: 2.2rem !important;
            font-weight: 800 !important;
            color: #ffffff !important;
            margin-bottom: 2px !important;
            text-shadow: 0 2px 10px rgba(0,0,0,0.5);
        }

        .hero-subtitle {
            font-size: 1rem !important;
            color: rgba(255, 255, 255, 0.85) !important;
            margin-bottom: 16px !important;
            font-weight: 500 !important;
        }

        /* Text Contrast Overrides */
        .stMarkdown, p, span, label, div[data-testid="stWidgetLabel"] {
            color: #ffffff !important;
            font-size: 0.95rem !important;
            font-weight: 600 !important;
        }

        .stTextInput input {
            background-color: rgba(0, 0, 0, 0.4) !important;
            color: #ffffff !important;
            border: 1px solid rgba(255, 255, 255, 0.2) !important;
            border-radius: 10px !important;
        }

        /* Radio Container Outer Box */
        div[data-testid="stRadio"] > div {
            background-color: rgba(0, 0, 0, 0.3) !important;
            padding: 6px 12px !important;
            border-radius: 10px !important;
            border: 1px solid rgba(255, 255, 255, 0.15) !important;
        }

        /* Clean transparent background for text labels */
        div[data-testid="stRadio"] label {
            background-color: transparent !important;
        }
        div[data-testid="stRadio"] label * {
            background-color: transparent !important;
            color: #ffffff !important;
        }

        /* Direct SVG / Radio Bullet targeting for Streamlit v1.63+ */
        div[data-testid="stRadio"] input[type="radio"]:checked + div {
            background-color: #9333ea !important;
            border-color: #9333ea !important;
        }
        div[data-testid="stRadio"] [data-baseweb="radio"] div {
            border-color: #9333ea !important;
        }

        div.stButton > button {
            background: linear-gradient(135deg, #9333ea 0%, #6b21a8 100%) !important;
            color: #ffffff !important;
            border-radius: 10px !important;
            font-size: 1rem !important;
            font-weight: 800 !important;
            border: none !important;
            box-shadow: 0 4px 15px rgba(147, 51, 234, 0.4) !important;
        }
        </style>

        <video autoplay muted loop playsinline id="bg-video">
            <source src="https://d8j0ntlcm91z4.cloudfront.net/user_38xzZboKViGWJOttwIXH07lWA1P/hf_20260826_124724_bc041163-d651-425f-aea3-2acc1efc2c96.mp4" type="video/mp4">
        </video>
    """, unsafe_allow_html=True)

    # Directly Position Sign In Form Over Motion Backdrop
    col_left, col_center, col_right = st.columns([0.15, 0.7, 0.15])

    with col_center:
        st.markdown('<div class="motion-login-card">', unsafe_allow_html=True)
        st.markdown('<div class="hero-title">Classmate AI Portal</div>', unsafe_allow_html=True)
        st.markdown('<div class="hero-subtitle">Summarize notes, generate quizzes, and solve doubts instantly.</div>', unsafe_allow_html=True)

        portal_type = st.radio("Select Portal Access:", ["Student Portal", "Admin Portal"], horizontal=True)

        if portal_type == "Student Portal":
            auth_mode = st.radio("Action:", ["Sign In", "Create Account"], horizontal=True)

            if auth_mode == "Sign In":
                login_email = st.text_input("Username / Email:", placeholder="student@example.com")
                login_pass = st.text_input("Password:", type="password")
                
                st.write("")
                if st.button("Login to Portal", use_container_width=True):
                    if login_email in st.session_state.student_db and st.session_state.student_db[login_email]["password"] == login_pass:
                        st.session_state.logged_in_user = st.session_state.student_db[login_email]
                        st.session_state.logged_in_email = login_email
                        st.session_state.user_role = "Student"
                        log_activity(st.session_state.logged_in_user["name"], login_email, "Portal Login", "Student logged into the Student Portal")
                        st.rerun()
                    else:
                        st.error("Invalid student email or password!")
            else:
                reg_name = st.text_input("Full Name:")
                reg_email = st.text_input("Register Email:")
                reg_pass = st.text_input("Register Password:", type="password")
                
                st.write("")
                if st.button("Create Account", use_container_width=True):
                    if reg_name and reg_email and reg_pass:
                        st.session_state.student_db[reg_email] = {"password": reg_pass, "name": reg_name}
                        save_json_data(STUDENT_DB_FILE, st.session_state.student_db)
                        log_activity(reg_name, reg_email, "Account Registration", "New student account created")
                        st.success("Account created successfully! Switch to 'Sign In'.")
                    else:
                        st.warning("Please fill in all registration fields.")

        else: # Admin Portal Selected
            st.info("🔒 Admin Access is restricted. Pre-authorized credentials required.")
            admin_email = st.text_input("Admin Email Address:", placeholder="admin@example.com")
            admin_pass = st.text_input("Admin Password:", type="password")
            
            st.write("")
            if st.button("Login to Admin Console", use_container_width=True):
                if admin_email in st.session_state.admin_db and st.session_state.admin_db[admin_email]["password"] == admin_pass:
                    st.session_state.logged_in_user = st.session_state.admin_db[admin_email]
                    st.session_state.logged_in_email = admin_email
                    st.session_state.user_role = "Admin"
                    log_activity(st.session_state.logged_in_user["name"], admin_email, "Admin Console Login", "Admin logged into the console")
                    st.rerun()
                else:
                    st.error("Access Denied! Invalid credentials or no Admin privileges.")

        st.markdown('</div>', unsafe_allow_html=True)

# ---------------------------------------------------------
# VIEW 1: STUDENT DASHBOARD
# ---------------------------------------------------------
elif st.session_state.user_role == "Student":
    user_name = st.session_state.logged_in_user["name"].strip()
    user_email = st.session_state.logged_in_email
    elapsed_minutes = int((time.time() - st.session_state.start_time) / 60)

    st.markdown(f"""
        <div class="welcome-card">
            <div class="welcome-title">Hi, {user_name}! 🎓📖</div>
            <div class="welcome-subtitle">Classmate AI Student Portal — Solve doubts with <b>Buddy</b>, summarize notes, and track real progress.</div>
        </div>
    """, unsafe_allow_html=True)

    col_a, col_b, col_c, col_d = st.columns(4)
    with col_a:
        st.markdown(f'<div class="dashboard-card">⏱️ <div class="card-title">Active Study Time</div><div class="card-subtext">{elapsed_minutes} Minutes</div></div>', unsafe_allow_html=True)
    with col_b:
        st.markdown(f'<div class="dashboard-card">📝 <div class="card-title">Notes Summarized</div><div class="card-subtext">{st.session_state.notes_processed} Processed</div></div>', unsafe_allow_html=True)
    with col_c:
        st.markdown(f'<div class="dashboard-card">❓ <div class="card-title">Quizzes Generated</div><div class="card-subtext">{st.session_state.quizzes_generated} Sets</div></div>', unsafe_allow_html=True)
    with col_d:
        st.markdown(f'<div class="dashboard-card">📅 <div class="card-title">Exams Scheduled</div><div class="card-subtext">{len(st.session_state.exam_schedule)} Scheduled</div></div>', unsafe_allow_html=True)

    st.write("")

    if nav == "🎓 AI Study Assistant":
        st.markdown("### 🛠️ AI Study Assistant Workspace")
        col1, col2 = st.columns([1, 1], gap="large")

        with col1:
            st.markdown('<div class="workspace-container">', unsafe_allow_html=True)
            st.markdown("#### 📝 Study Material Input")
            
            feature = st.selectbox(
                "Select Action:",
                ["Summarize Notes", "Generate Practice Quiz", "Explain Concept Simply"]
            )

            input_mode = st.radio("Input Method:", ["Paste Text", "Upload Document/Image"])
            extracted_text = ""
            uploaded_file = None
            file_bytes_payload = None

            if input_mode == "Paste Text":
                extracted_text = st.text_area("Enter study material or notes:", height=180, placeholder="Paste study notes or text on any topic...")
            else:
                uploaded_file = st.file_uploader("Upload File (PDF, PPT, Excel, Image):", type=["pdf", "pptx", "xlsx", "xls", "csv", "png", "jpg", "jpeg", "txt"])
                if uploaded_file is not None:
                    file_bytes_payload = uploaded_file.getvalue()
                    file_type = uploaded_file.name.split(".")[-1].lower()
                    try:
                        if file_type == "pdf":
                            reader = pypdf.PdfReader(uploaded_file)
                            for page in reader.pages:
                                text = page.extract_text()
                                if text: extracted_text += text + "\n"
                        elif file_type == "pptx":
                            prs = Presentation(uploaded_file)
                            for slide in prs.slides:
                                for shape in slide.shapes:
                                    if hasattr(shape, "text"): extracted_text += shape.text + "\n"
                        elif file_type in ["xlsx", "xls", "csv"]:
                            df = pd.read_csv(uploaded_file) if file_type == "csv" else pd.read_excel(uploaded_file)
                            extracted_text = df.to_string()
                        elif file_type == "txt":
                            extracted_text = uploaded_file.read().decode("utf-8")
                        elif file_type in ["png", "jpg", "jpeg"]:
                            st.image(uploaded_file, caption="Uploaded Image", use_container_width=True)
                            
                        st.success(f"Parsed: {uploaded_file.name}")
                    except Exception as e:
                        st.error(f"Error parsing file: {str(e)}")

            generate_btn = st.button("🚀 Process with AI", use_container_width=True)
            st.markdown('</div>', unsafe_allow_html=True)

        with col2:
            st.markdown('<div class="workspace-container">', unsafe_allow_html=True)
            st.markdown("#### 💡 AI Output")
            
            if generate_btn:
                if input_mode == "Paste Text" and not extracted_text.strip():
                    st.warning("⚠️ Please provide text before submitting!")
                elif input_mode == "Upload Document/Image" and uploaded_file is None:
                    st.warning("⚠️ Please upload a file first!")
                else:
                    if feature == "Summarize Notes":
                        instruction = "Summarize the following study material into clear, structured bullet points with headers:"
                        st.session_state.notes_processed += 1
                    elif feature == "Generate Practice Quiz":
                        instruction = "Create 3 practice multiple-choice questions (A-D) with correct answers and explanations based on this study material:"
                        st.session_state.quizzes_generated += 1
                    else:
                        instruction = "Explain the main concepts in this study material in simple terms with examples:"
                        st.session_state.concepts_explained += 1

                    try:
                        with st.spinner("Classmate AI processing..."):
                            if input_mode == "Upload Document/Image":
                                ext = uploaded_file.name.split(".")[-1].lower()
                                if ext in ["png", "jpg", "jpeg"]:
                                    image_input = Image.open(io.BytesIO(file_bytes_payload))
                                    contents = [image_input, instruction]
                                elif ext == "pdf" and not extracted_text.strip():
                                    pdf_part = types.Part.from_bytes(
                                        data=file_bytes_payload,
                                        mime_type="application/pdf"
                                    )
                                    contents = [pdf_part, instruction]
                                else:
                                    contents = f"{instruction}\n\n{extracted_text}"
                            else:
                                contents = f"{instruction}\n\n{extracted_text}"

                            res = query_gemini(contents)
                            st.markdown(res)
                            log_activity(user_name, user_email, f"Used AI Assistant: {feature}", f"Mode: {input_mode}")
                    except Exception as e:
                        st.error(f"❌ {str(e)}")
            else:
                st.info("AI generated output will appear here.")
            st.markdown('</div>', unsafe_allow_html=True)

    elif nav == "📋 Custom Study Plan":
        st.subheader("📋 Custom Study Plan (To-Do List)")
        col_td1, col_td2 = st.columns([1, 1], gap="large")

        with col_td1:
            st.markdown('<div class="workspace-container">', unsafe_allow_html=True)
            st.markdown("### ➕ Add New Task")
            new_task = st.text_input("Enter study goal or topic:", placeholder="e.g. Read Physics Chapter 3")
            if st.button("Add to Study Plan", use_container_width=True):
                if new_task.strip():
                    st.session_state.todo_list.append({"task": new_task, "done": False})
                    log_activity(user_name, user_email, "Added Study Plan Task", f"Task: {new_task}")
                    st.success("Task added!")
                    st.rerun()
                else:
                    st.warning("Please enter a task name.")
            st.markdown('</div>', unsafe_allow_html=True)

        with col_td2:
            st.markdown('<div class="workspace-container">', unsafe_allow_html=True)
            st.markdown("### 📌 Active Tasks")
            if st.session_state.todo_list:
                for idx, item in enumerate(st.session_state.todo_list):
                    col_chk, col_btn = st.columns([4, 1])
                    with col_chk:
                        done = st.checkbox(item["task"], value=item["done"], key=f"task_{idx}")
                        if done != item["done"]:
                            st.session_state.todo_list[idx]["done"] = done
                            status_str = "Completed" if done else "Marked Incomplete"
                            log_activity(user_name, user_email, f"Updated Study Task ({status_str})", f"Task: {item['task']}")
                    with col_btn:
                        if st.button("❌", key=f"del_task_{idx}"):
                            removed_item = st.session_state.todo_list.pop(idx)
                            log_activity(user_name, user_email, "Deleted Study Task", f"Task: {removed_item['task']}")
                            st.rerun()
            else:
                st.info("No study tasks added yet!")
            st.markdown('</div>', unsafe_allow_html=True)

    elif nav == "💬 Buddy (AI Chatbot)":
        st.subheader("💬 Buddy — Universal Doubt Clearing Chatbot")
        st.write("Ask Buddy any doubt regarding **any subject or topic** (Math, Science, History, Coding, Languages, etc.)!")

        for message in st.session_state.messages:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])

        active_prompt = st.chat_input("Ask Buddy a question about any subject...")

        if active_prompt:
            st.session_state.messages.append({"role": "user", "content": active_prompt})
            with st.chat_message("user"):
                st.markdown(active_prompt)

            with st.chat_message("assistant"):
                try:
                    with st.spinner("Buddy is answering..."):
                        chat_prompt = f"You are Buddy, a friendly expert tutor for all academic subjects. Answer the student's question clearly and step-by-step.\n\nStudent Question: {active_prompt}"
                        reply_text = query_gemini(chat_prompt)
                        st.markdown(reply_text)
                        st.session_state.messages.append({"role": "assistant", "content": reply_text})
                        log_activity(user_name, user_email, "Asked Buddy (Chatbot)", f"Prompt: {active_prompt[:60]}...")
                except Exception as e:
                    st.error(f"Chatbot Error: {str(e)}")

    elif nav == "📅 Exam Schedule":
        st.subheader("📅 Dynamic Exam Schedule Manager")
        col_sch1, col_sch2 = st.columns([1, 1], gap="large")
        
        with col_sch1:
            st.markdown('<div class="workspace-container">', unsafe_allow_html=True)
            st.markdown("### ➕ Add Exam / Task")
            new_subject = st.text_input("Exam Name:", placeholder="e.g. Physics Midterm")
            new_date = st.date_input("Exam Date:")
            
            if st.button("Save Exam", use_container_width=True):
                if new_subject.strip():
                    st.session_state.exam_schedule.append({"subject": new_subject, "date": str(new_date)})
                    log_activity(user_name, user_email, "Added Scheduled Exam", f"Exam: {new_subject} on {new_date}")
                    st.success(f"Added **{new_subject}** on `{new_date}`!")
                else:
                    st.warning("Enter an exam name.")
            st.markdown('</div>', unsafe_allow_html=True)

        with col_sch2:
            st.markdown('<div class="workspace-container">', unsafe_allow_html=True)
            st.markdown("### 📋 Scheduled Exams")
            if st.session_state.exam_schedule:
                for idx, item in enumerate(st.session_state.exam_schedule):
                    st.markdown(f"**{idx + 1}. {item['subject']}** — 🗓️ `{item['date']}`")
            else:
                st.info("No exams scheduled.")
            st.markdown('</div>', unsafe_allow_html=True)

    else:
        st.subheader("📊 Real Session Study Analytics")
        st.markdown("Real-time activity stats generated from your current study session:")
        
        col_m1, col_m2, col_m3 = st.columns(3)
        col_m1.metric("Active Session Time", f"{elapsed_minutes} mins")
        col_m2.metric("Notes Summarized", f"{st.session_state.notes_processed}")
        col_m3.metric("Quizzes Completed", f"{st.session_state.quizzes_generated}")

        st.divider()
        st.markdown("**Real Activity Breakdown**")
        activity_data = pd.DataFrame({
            "Activity": ["Summaries", "Quizzes", "Concepts Explained"],
            "Count": [st.session_state.notes_processed, st.session_state.quizzes_generated, st.session_state.concepts_explained]
        })
        st.bar_chart(activity_data.set_index("Activity"))

# ---------------------------------------------------------
# VIEW 2: ADMIN DASHBOARD
# ---------------------------------------------------------
else:
    admin_name = st.session_state.logged_in_user["name"].strip()

    st.markdown(f"""
        <div class="welcome-card">
            <div class="welcome-title">🛡️ Admin Control Console</div>
            <div class="welcome-subtitle">Welcome, <b>{admin_name}</b>. Manage student accounts, monitor live activity logs, grant admin privileges, and manage exams.</div>
        </div>
    """, unsafe_allow_html=True)

    admin_tabs = st.tabs([
        "👥 Student Accounts & Access Control", 
        "📊 Live Student Usage Logs", 
        "🛡️ Active Admin List", 
        "📅 System Exam Manager"
    ])

    # Tab 1: Access Control
    with admin_tabs[0]:
        st.subheader("Manage Registered Students & Grant Admin Access")
        
        if st.session_state.student_db:
            for email, info in list(st.session_state.student_db.items()):
                col_u1, col_u2 = st.columns([3, 1])
                with col_u1:
                    st.write(f"👤 **{info['name'].strip()}** (`{email}`)")
                with col_u2:
                    is_admin = email in st.session_state.admin_db
                    if is_admin:
                        st.success("Admin Access Granted")
                    else:
                        if st.button(f"Grant Admin Access", key=f"grant_{email}", use_container_width=True):
                            st.session_state.admin_db[email] = {
                                "password": info["password"],
                                "name": f"{info['name'].strip()} (Admin)"
                            }
                            st.success(f"Granted Admin privileges to {info['name'].strip()}!")
                            st.rerun()
        else:
            st.info("No registered student accounts yet.")

    # Tab 2: Activity Logs View
    with admin_tabs[1]:
        st.subheader("📊 Live Student Usage & Activity Logs")
        st.caption("Track logins, tool usage, chatbot prompts, and study tasks in real-time.")

        if st.session_state.user_activity_log:
            df_logs = pd.DataFrame(st.session_state.user_activity_log)
            st.dataframe(df_logs, use_container_width=True)
        else:
            st.info("No student activity logged yet in this session.")

    # Tab 3: Admin List
    with admin_tabs[2]:
        st.subheader("Active Authorized Admins")
        admins_list = [{"Admin Name": info["name"], "Email": email} for email, info in st.session_state.admin_db.items()]
        st.dataframe(pd.DataFrame(admins_list), use_container_width=True)

    # Tab 4: System Exams
    with admin_tabs[3]:
        st.subheader("Active System Exams")
        if st.session_state.exam_schedule:
            for idx, item in enumerate(st.session_state.exam_schedule):
                col_e1, col_e2 = st.columns([3, 1])
                with col_e1: st.write(f"**{idx+1}. {item['subject']}** (`{item['date']}`)")
                with col_e2:
                    if st.button(f"Delete #{idx+1}", key=f"del_{idx}", use_container_width=True):
                        st.session_state.exam_schedule.pop(idx)
                        st.rerun()
        else:
            st.info("No exams in database.")
